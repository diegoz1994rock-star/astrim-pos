"""Prueba de humo de las opciones nuevas y opcionales de
`export_table_to_excel` (`header_fill_hex`/`autofit`, agregadas para
Ganancias) — el comportamiento por defecto (sin esos parámetros) debe
seguir generando el mismo archivo que antes."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from pos.modules.reports.infrastructure.exporter import export_table_to_excel


def test_default_call_still_works_without_new_params(tmp_path: Path) -> None:
    file_path = tmp_path / "reporte.xlsx"

    export_table_to_excel(
        title="Reporte", headers=["A", "B"], rows=[["1", "2"]], file_path=file_path
    )

    assert file_path.exists()


def test_header_fill_and_autofit_apply_when_requested(tmp_path: Path) -> None:
    file_path = tmp_path / "ganancias.xlsx"

    export_table_to_excel(
        title="Ganancias",
        headers=["Código", "Producto muy largo de nombre"],
        rows=[["SKU-1", "Martillo"]],
        file_path=file_path,
        header_fill_hex="#2F6FED",
        autofit=True,
    )

    workbook = load_workbook(file_path)
    sheet = workbook.active
    header_cell = sheet.cell(row=1, column=2)
    assert header_cell.fill.fgColor.rgb.endswith("2F6FED")
    assert sheet.column_dimensions["B"].width > 10
