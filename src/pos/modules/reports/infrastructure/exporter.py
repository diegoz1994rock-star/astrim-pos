"""Exportación genérica de tablas de reporte a PDF y Excel.

Genérico a propósito: todo reporte del módulo (ventas, inventario, caja,
productos, etc.) se reduce a una lista de encabezados + filas de texto antes
de exportar, así este exportador no necesita conocer el dominio de cada
reporte (PROJECT_SPEC.md, "REPORTES": todos exportables a PDF y Excel).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def export_table_to_pdf(
    *, title: str, headers: list[str], rows: list[list[str]], file_path: Path
) -> None:
    file_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(file_path), pagesize=letter)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    table_data = [headers, *rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F6FED")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F6F8")]),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)


def export_table_to_excel(
    *,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    file_path: Path,
    header_fill_hex: str | None = None,
    autofit: bool = False,
) -> None:
    """`header_fill_hex`/`autofit` son opcionales (default = comportamiento
    de siempre, sin cambios para los reportes existentes) — agregados
    para el formato "profesional" que pide Ganancias (encabezado con
    color, columnas autoajustadas) sin duplicar esta función."""
    file_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title[:31] or "Reporte"

    header_fill = None
    if header_fill_hex:
        header_fill = PatternFill("solid", fgColor=header_fill_hex.lstrip("#"))
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF" if header_fill else None)
        if header_fill:
            cell.fill = header_fill

    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=col_index, value=value)

    if autofit:
        for col_index, header in enumerate(headers, start=1):
            widest = max(
                [len(str(header))] + [len(str(row[col_index - 1])) for row in rows if row],
                default=len(str(header)),
            )
            sheet.column_dimensions[get_column_letter(col_index)].width = min(widest + 2, 60)

    workbook.save(str(file_path))
